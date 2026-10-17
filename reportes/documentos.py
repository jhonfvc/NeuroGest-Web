"""HU-19 / HU-20: generación de la receta en PDF y exportación de reportes (PDF / Excel)."""
from io import BytesIO

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, A5
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INDIGO = colors.HexColor('#4338ca')
GRIS = colors.HexColor('#64748b')

_estilos = getSampleStyleSheet()
TITULO = ParagraphStyle('titulo', parent=_estilos['Title'], fontSize=15, textColor=INDIGO, spaceAfter=0)
SUBTITULO = ParagraphStyle('subtitulo', parent=_estilos['Normal'], fontSize=8, textColor=GRIS, alignment=TA_CENTER)
NORMAL = ParagraphStyle('normal', parent=_estilos['Normal'], fontSize=9, leading=12)
PEQUENO = ParagraphStyle('pequeno', parent=NORMAL, fontSize=8, textColor=GRIS)
SECCION = ParagraphStyle('seccion', parent=NORMAL, fontName='Helvetica-Bold', fontSize=10, textColor=INDIGO,
                         spaceBefore=6, spaceAfter=3)


def _membrete():
    c = settings.CONSULTORIO
    datos = ' · '.join(x for x in (c['direccion'], c['telefono']) if x)
    return [Paragraph(c['nombre'], TITULO), Paragraph(datos, SUBTITULO), Spacer(1, 3 * mm),
            HRFlowable(width='100%', color=INDIGO, thickness=1), Spacer(1, 3 * mm)]


def _esc(texto):
    return (str(texto or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')).replace('\n', '<br/>')


def receta_pdf(atencion):
    """HU-19: receta con membrete, datos del paciente, diagnóstico CIE-10, prescripción y área de firma."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A5, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm,
                            title=f'Receta {atencion.paciente.numero_historia}', author=settings.CONSULTORIO['nombre'])
    p = atencion.paciente
    medico = atencion.medico.get_full_name() or atencion.medico.username
    fecha = timezone.localtime(atencion.fecha)

    historia = _membrete()
    historia.append(Paragraph('RECETA MÉDICA', ParagraphStyle('r', parent=SECCION, alignment=TA_CENTER, fontSize=12)))
    datos = [
        [Paragraph(f'<b>Paciente:</b> {_esc(p.nombre_completo)}', NORMAL), Paragraph(f'<b>Fecha:</b> {fecha:%d/%m/%Y}', NORMAL)],
        [Paragraph(f'<b>DNI:</b> {p.dni}', NORMAL), Paragraph(f'<b>N° Historia:</b> {p.numero_historia}', NORMAL)],
        [Paragraph(f'<b>Edad:</b> {p.edad} años' if p.edad is not None else '', NORMAL),
         Paragraph(f'<b>Peso:</b> {atencion.triaje.peso} kg' if hasattr(atencion, 'triaje') else '', NORMAL)],
    ]
    tabla = Table(datos, colWidths=['62%', '38%'])
    tabla.setStyle(TableStyle([('LEFTPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 1)]))
    historia.append(tabla)

    historia.append(Paragraph('Diagnóstico (CIE-10)', SECCION))
    for d in atencion.diagnosticos.select_related('cie10'):
        historia.append(Paragraph(f'<b>{d.cie10.codigo}</b> {_esc(d.cie10.descripcion)}{" (principal)" if d.principal else ""}', NORMAL))

    historia.append(Paragraph('Prescripción', SECCION))
    prescripciones = list(atencion.prescripciones.all())
    if prescripciones:
        for i, rx in enumerate(prescripciones, 1):
            nombre = f'{i}. <b>{_esc(rx.medicamento)}</b>' + (f' {_esc(rx.presentacion)}' if rx.presentacion else '')
            historia.append(Paragraph(nombre, NORMAL))
            detalle = rx.posologia + (f'. {_esc(rx.indicaciones)}' if rx.indicaciones else '')
            historia.append(Paragraph(detalle, ParagraphStyle('rx', parent=NORMAL, leftIndent=10, textColor=colors.HexColor('#334155'))))
            historia.append(Spacer(1, 2 * mm))
    else:
        historia.append(Paragraph('Sin medicamentos prescritos.', PEQUENO))

    if atencion.plan:
        historia.append(Paragraph('Indicaciones', SECCION))
        historia.append(Paragraph(_esc(atencion.plan), NORMAL))

    historia += [Spacer(1, 18 * mm), HRFlowable(width='45%', color=colors.black, thickness=0.6, hAlign='CENTER'),
                 Paragraph(f'Dr(a). {_esc(medico)}', ParagraphStyle('f', parent=NORMAL, alignment=TA_CENTER)),
                 Paragraph('Médico Neurólogo · Firma y sello', ParagraphStyle('f2', parent=PEQUENO, alignment=TA_CENTER))]
    doc.build(historia)
    return buffer.getvalue()


FILAS_REPORTE = ('Atendidas', 'Canceladas', 'Ausentes', 'Pendientes')
CLAVES = ('atendidas', 'canceladas', 'ausentes', 'pendientes')


def _titulo_reporte(resumen):
    medico = f" · Dr(a). {resumen.medico.get_full_name() or resumen.medico.username}" if resumen.medico else ''
    return f"Reporte de atenciones del {resumen.inicio:%d/%m/%Y} al {resumen.fin:%d/%m/%Y}{medico}"


def reporte_pdf(resumen):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=_titulo_reporte(resumen))
    estilo_tabla = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), INDIGO), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (1, 0), (-1, -1), 'RIGHT'), ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ])
    historia = _membrete() + [Paragraph(_titulo_reporte(resumen), SECCION)]

    totales = [['Indicador', 'Total']] + [[n, resumen.totales[k]] for n, k in zip(FILAS_REPORTE, CLAVES)]
    totales.append(['Total de citas', resumen.total])
    if resumen.tasa_asistencia is not None:
        totales.append(['Tasa de asistencia', f'{resumen.tasa_asistencia} %'])
    t = Table(totales, colWidths=[70 * mm, 30 * mm], hAlign='LEFT')
    t.setStyle(estilo_tabla)
    historia += [t, Spacer(1, 5 * mm)]

    if resumen.por_medico:
        historia.append(Paragraph('Por médico', SECCION))
        filas = [['Médico', *FILAS_REPORTE]] + [
            [f"{m['medico__first_name']} {m['medico__last_name']}".strip() or m['medico__username'], *(m[k] for k in CLAVES)]
            for m in resumen.por_medico]
        t = Table(filas, colWidths=[60 * mm] + [28 * mm] * 4, hAlign='LEFT')
        t.setStyle(estilo_tabla)
        historia += [t, Spacer(1, 5 * mm)]

    historia.append(Paragraph('Por día', SECCION))
    filas = [['Fecha', *FILAS_REPORTE]] + [[f"{d['fecha']:%d/%m/%Y}", *(d[k] for k in CLAVES)]
                                            for d in resumen.por_dia if any(d[k] for k in CLAVES)]
    if len(filas) == 1:
        historia.append(Paragraph('No hay citas en el rango seleccionado.', PEQUENO))
    else:
        t = Table(filas, colWidths=[40 * mm] + [28 * mm] * 4, hAlign='LEFT', repeatRows=1)
        t.setStyle(estilo_tabla)
        historia.append(t)
    historia.append(Spacer(1, 5 * mm))
    historia.append(Paragraph(f'Generado el {timezone.localtime():%d/%m/%Y %H:%M} · NeuroGest', PEQUENO))
    doc.build(historia)
    return buffer.getvalue()


def reporte_excel(resumen):
    libro = Workbook()
    negrita = Font(bold=True, color='FFFFFF')
    relleno = PatternFill('solid', fgColor='4338CA')

    def encabezado(hoja, fila, valores):
        for col, valor in enumerate(valores, 1):
            celda = hoja.cell(row=fila, column=col, value=valor)
            celda.font, celda.fill, celda.alignment = negrita, relleno, Alignment(horizontal='center')

    hoja = libro.active
    hoja.title = 'Resumen'
    hoja['A1'] = _titulo_reporte(resumen)
    hoja['A1'].font = Font(bold=True, size=13)
    encabezado(hoja, 3, ['Indicador', 'Total'])
    for i, (nombre, clave) in enumerate(zip(FILAS_REPORTE, CLAVES), 4):
        hoja.cell(row=i, column=1, value=nombre)
        hoja.cell(row=i, column=2, value=resumen.totales[clave])
    hoja.cell(row=8, column=1, value='Total de citas').font = Font(bold=True)
    hoja.cell(row=8, column=2, value=resumen.total).font = Font(bold=True)
    if resumen.tasa_asistencia is not None:
        hoja.cell(row=9, column=1, value='Tasa de asistencia (%)')
        hoja.cell(row=9, column=2, value=resumen.tasa_asistencia)
    hoja.column_dimensions['A'].width = 28

    dias = libro.create_sheet('Por día')
    encabezado(dias, 1, ['Fecha', *FILAS_REPORTE])
    for i, d in enumerate(resumen.por_dia, 2):
        dias.cell(row=i, column=1, value=d['fecha']).number_format = 'DD/MM/YYYY'
        for col, clave in enumerate(CLAVES, 2):
            dias.cell(row=i, column=col, value=d[clave])
    dias.column_dimensions['A'].width = 14

    medicos = libro.create_sheet('Por médico')
    encabezado(medicos, 1, ['Médico', *FILAS_REPORTE])
    for i, m in enumerate(resumen.por_medico, 2):
        medicos.cell(row=i, column=1, value=f"{m['medico__first_name']} {m['medico__last_name']}".strip() or m['medico__username'])
        for col, clave in enumerate(CLAVES, 2):
            medicos.cell(row=i, column=col, value=m[clave])
    medicos.column_dimensions['A'].width = 30

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
