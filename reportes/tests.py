from datetime import time, timedelta
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from admision.models import Cita, Paciente
from hce.models import Atencion, Cie10, Diagnostico, Prescripcion, Triaje
from seguridad.models import Usuario

from .servicios import generar_resumen

CLAVE = 'ClaveSegura#2026'


def usuario(username, rol, **extra):
    return Usuario.objects.create_user(username, email=f'{username}@neurogest.pe', password=CLAVE, rol=rol, **extra)


class BaseReportes(TestCase):
    def setUp(self):
        self.admin = usuario('admin', Usuario.ADMIN)
        self.medico = usuario('medico', Usuario.MEDICO, first_name='Miguel', last_name='Salazar')
        self.medico2 = usuario('medico2', Usuario.MEDICO, first_name='Lucía', last_name='Paredes')
        self.paciente = Paciente.objects.create(dni='12345678', nombres='Rosa', apellido_paterno='Quispe')
        self.hoy = timezone.localdate()
        self.ayer = self.hoy - timedelta(days=1)

    def cita(self, fecha, hora, estado, medico=None):
        return Cita.objects.create(paciente=self.paciente, medico=medico or self.medico, fecha=fecha, hora=time(hora), estado=estado)


class ResumenTests(BaseReportes):
    """HU-20: conteo de atendidas, canceladas y ausentes."""

    def setUp(self):
        super().setUp()
        self.cita(self.ayer, 9, Cita.ATENDIDA)
        self.cita(self.ayer, 10, Cita.CANCELADA)
        self.cita(self.ayer, 11, Cita.AGENDADA)          # ausente: pasó y nunca llegó
        self.cita(self.hoy, 9, Cita.AGENDADA)            # pendiente
        self.cita(self.ayer, 9, Cita.ATENDIDA, self.medico2)
        self.cita(self.hoy - timedelta(days=40), 9, Cita.ATENDIDA)  # fuera de rango

    def test_totales(self):
        r = generar_resumen(self.ayer, self.hoy)
        self.assertEqual(r.totales, {'atendidas': 2, 'canceladas': 1, 'ausentes': 1, 'pendientes': 1})
        self.assertEqual(r.total, 5)
        self.assertEqual(r.tasa_asistencia, 66.7)

    def test_filtro_por_medico(self):
        r = generar_resumen(self.ayer, self.hoy, self.medico2)
        self.assertEqual(r.totales['atendidas'], 1)
        self.assertEqual(r.total, 1)

    def test_serie_por_dia_incluye_dias_sin_citas(self):
        r = generar_resumen(self.ayer - timedelta(days=2), self.hoy)
        self.assertEqual(len(r.por_dia), 4)
        self.assertEqual(r.por_dia[0]['atendidas'], 0)
        self.assertEqual(r.por_dia[2]['atendidas'], 2)

    def test_por_medico(self):
        r = generar_resumen(self.ayer, self.hoy)
        self.assertEqual([m['medico__last_name'] for m in r.por_medico], ['Paredes', 'Salazar'])

    def test_rango_invertido(self):
        r = generar_resumen(self.hoy, self.ayer)
        self.assertEqual((r.inicio, r.fin), (self.ayer, self.hoy))


class PanelTests(BaseReportes):
    def setUp(self):
        super().setUp()
        self.cita(self.ayer, 9, Cita.ATENDIDA)
        self.rango = {'inicio': self.ayer.isoformat(), 'fin': self.hoy.isoformat()}

    def test_panel(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('reportes'), self.rango)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['resumen'].totales['atendidas'], 1)

    def test_exportar_excel(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('reporte_excel'), self.rango)
        self.assertIn('attachment', resp['Content-Disposition'])
        libro = load_workbook(BytesIO(resp.content))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Por día', 'Por médico'])
        self.assertEqual(libro['Resumen']['B4'].value, 1)

    def test_exportar_pdf(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('reporte_pdf'), self.rango)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_solo_administrador(self):
        self.client.force_login(self.medico)
        self.assertEqual(self.client.get(reverse('reportes')).status_code, 403)
        self.assertEqual(self.client.get(reverse('reporte_excel')).status_code, 403)


class RecetaPdfTests(BaseReportes):
    """HU-19: receta en PDF con membrete, paciente, diagnóstico y prescripción."""

    def setUp(self):
        super().setUp()
        self.atencion = Atencion.objects.create(paciente=self.paciente, medico=self.medico,
                                                motivo_consulta='Cefalea', evolucion='Estable', plan='Reposo & hidratación')
        Triaje.objects.create(atencion=self.atencion, presion_sistolica=120, presion_diastolica=80, peso=70, talla=170)
        Diagnostico.objects.create(atencion=self.atencion, cie10=Cie10.objects.get(codigo='G43.9'), principal=True)
        Prescripcion.objects.create(atencion=self.atencion, medicamento='Naproxeno', dosis=550, unidad='mg',
                                    frecuencia_horas=12, duracion_dias=5)

    def test_genera_pdf(self):
        self.client.force_login(self.medico)
        resp = self.client.get(reverse('receta_pdf', args=[self.atencion.pk]))
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))
        self.assertIn('inline', resp['Content-Disposition'])

    def test_descarga_en_un_clic(self):
        self.client.force_login(self.medico)
        resp = self.client.get(reverse('receta_pdf', args=[self.atencion.pk]), {'descargar': 1})
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertIn(self.paciente.numero_historia, resp['Content-Disposition'])

    def test_recepcion_no_accede(self):
        self.client.force_login(usuario('recep', Usuario.RECEPCION))
        self.assertEqual(self.client.get(reverse('receta_pdf', args=[self.atencion.pk])).status_code, 403)
